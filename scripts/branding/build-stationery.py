"""Build editable office templates; dimensions are custom, not a label-stock preset."""

from datetime import datetime, timezone
from io import BytesIO
import json
from pathlib import Path
import uuid
from zipfile import ZipFile, ZIP_DEFLATED, ZipInfo

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree
import qrcode

ROOT = Path(__file__).resolve().parents[2]
BRAND = ROOT / "branding"
TOKENS = json.loads((BRAND / "tokens.json").read_text(encoding="utf-8"))
OUT = BRAND / "office"
BLUE, INK, MUTED = (TOKENS["color"][key].lstrip("#") for key in ["blue", "ink", "muted"])
FONT, MONO = TOKENS["font"]["brand"], TOKENS["font"]["code"]
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}


def element(name, **attrs):
    value = OxmlElement(name)
    for key, item in attrs.items():
        value.set(qn(key), str(item))
    return value


def run(p, text, size=11, bold=False, mono=False, ink=INK):
    value = p.add_run(text)
    value.font.name = MONO if mono else FONT
    value.font.size, value.bold, value.font.color.rgb = Pt(size), bold, RGBColor.from_string(ink)
    value._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), value.font.name)
    return value


def paragraph(parent, text="", size=11, before=0, after=8, **kwargs):
    p = parent.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        run(p, text, size, **kwargs)
    return p


def link(p, label, url, size=9):
    relationship = p.part.relate_to(url, NS["r"] + "/hyperlink", is_external=True)
    h = element("w:hyperlink", **{"r:id": relationship})
    r = element("w:r")
    props = element("w:rPr")
    props.append(element("w:rFonts", **{"w:ascii": MONO, "w:hAnsi": MONO}))
    props.append(element("w:sz", **{"w:val": int(size * 2)}))
    props.append(element("w:color", **{"w:val": BLUE}))
    r.append(props)
    text = element("w:t")
    text.text = label
    r.append(text)
    h.append(r)
    p._p.append(h)


def qr():
    image = qrcode.make("https://" + TOKENS["domain"], border=4)
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


def base(title, top=45, side=22):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(top), Mm(29)
    section.left_margin = section.right_margin = Mm(side)
    section.header_distance, section.footer_distance = Mm(14), Mm(12)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(8)
    for name, size in [("Title", 24), ("Heading 1", 16), ("Heading 2", 12)]:
        style = doc.styles[name]
        style.font.name, style.font.size = FONT, Pt(size)
        style.font.bold, style.font.color.rgb = True, RGBColor(0, 0, 0)
        for attr in list(style.element.rPr.rFonts.attrib):
            if 'Theme' in attr:
                del style.element.rPr.rFonts.attrib[attr]
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    lang = element("w:lang", **{"w:val": "en-GB"})
    normal.element.get_or_add_rPr().append(lang)
    settings = doc.settings.element
    for setting in settings.xpath('./w:compat/w:compatSetting'):
        if setting.get(qn('w:name')) == 'compatibilityMode':
            setting.set(qn('w:val'), '15')
    settings.append(element("w:embedTrueTypeFonts"))
    settings.append(element("w:saveSubsetFonts", **{"w:val": "false"}))
    doc.core_properties.title = title
    doc.core_properties.subject = "FNLLA brand edition " + TOKENS["edition"]
    doc.core_properties.author = TOKENS["maintainer"]
    doc.core_properties.last_modified_by = TOKENS["maintainer"]
    doc.core_properties.comments = ""
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 9, 7, tzinfo=timezone.utc)
    return doc


def identity(doc):
    section = doc.sections[0]
    p = section.header.paragraphs[0]
    p.paragraph_format.tab_stops.clear_all()
    p.paragraph_format.tab_stops.add_tab_stop(Mm(166), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run().add_picture(str(BRAND / "assets/logo/wordmark-transparent.png"), width=Mm(49))
    p.add_run("\t").add_picture(str(BRAND / "assets/patterns/binary-field.png"), width=Mm(30))
    footer = section.footer
    p = footer.paragraphs[0]
    p.paragraph_format.space_after = Pt(5)
    link(p, TOKENS["domain"], "https://" + TOKENS["domain"], 10)
    run(p, "   /   ", 9, ink=MUTED)
    link(p, TOKENS["maintainer_email"], "mailto:" + TOKENS["maintainer_email"], 9)
    p = paragraph(footer, "Created & maintained by " + TOKENS["maintainer"] + "   |   ", 9, after=3)
    link(p, TOKENS["maintainer_domain"], "https://" + TOKENS["maintainer_domain"], 9)
    paragraph(footer, "Lead Developer / Product Manager - " + TOKENS["lead"], 8.5, after=0, ink=MUTED)


def letter(cover=False):
    doc = base("FNLLA cover letter" if cover else "FNLLA letterhead")
    identity(doc)
    paragraph(doc, "[Day Month Year]", 9, mono=True, after=15)
    paragraph(doc, "[Recipient name]\n[Organisation]\n[Postal address]", 11, after=22)
    doc.add_paragraph("Project correspondence" if cover else "Letter subject", "Title")
    paragraph(doc, "[Project name or reference]", 9, mono=True, ink=MUTED, after=22)
    paragraph(doc, "Dear [Recipient name]", after=15)
    if cover:
        paragraph(doc, "Please find enclosed [document or proposal title] for [project or purpose]. "
                  "It sets out [scope of the submission] and the decision requested from your team.")
        paragraph(doc, "[Summarise the key point, relevant context and any confirmed constraints. "
                  "Replace this paragraph with the details of your submission.]")
        paragraph(doc, "Please [review, approve or respond with the requested action] by [agreed date]. "
                  "For questions, contact us using the details below.")
    else:
        paragraph(doc, "[State the purpose of this letter and the action or decision requested.]")
        paragraph(doc, "[Provide the relevant context, confirmed facts and supporting details. "
                  "Add paragraphs as needed; the letterhead repeats on subsequent pages.]")
        paragraph(doc, "[Close with the next step, responsible contact and any agreed date.]")
    paragraph(doc, "Yours sincerely", before=12, after=24)
    paragraph(doc, TOKENS["lead"], 11, bold=True, after=3)
    paragraph(doc, "Lead Developer / Product Manager", 10, after=3)
    paragraph(doc, TOKENS["maintainer"], 10, after=12)
    if cover:
        paragraph(doc, "Enclosures  [Document title and revision]", 9, ink=MUTED)
    return doc


def grid(doc, width, height, rows, gap):
    table = doc.add_table(rows=rows * 2 - 1, cols=3)
    table.autofit = False
    for col, mm in zip(table.columns, [width, gap, width]):
        col.width = Mm(mm)
    for row_index, row in enumerate(table.rows):
        row.height = Mm(gap if row_index % 2 else height)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        row._tr.get_or_add_trPr().append(element("w:cantSplit"))
        for col_index, cell in enumerate(row.cells):
            active = row_index % 2 == 0 and col_index != 1
            cell.width = Mm(gap if col_index == 1 else width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = element("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                margins.append(element("w:" + side, **{"w:w": 190 if active and side in ['left', 'right'] else 0, "w:type": "dxa"}))
            props.append(margins)
            borders = element("w:tcBorders")
            for edge in ["top", "left", "bottom", "right"]:
                borders.append(element("w:" + edge, **{"w:val": "single" if row_index % 2 == 0 and col_index != 1 else "nil", "w:sz": 3, "w:color": "D9D9D9"}))
            props.append(borders)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1
            run(p, "", 1)
            if not active:
                p.paragraph_format.line_spacing = Pt(1)
    return [table.cell(r * 2, c) for r in range(rows) for c in [0, 2]]


def cards():
    doc = base("FNLLA business cards", top=26, side=15)
    doc.sections[0].bottom_margin = Mm(15)
    for cell in grid(doc, 85, 55, 4, 5):
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(77), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run().add_picture(str(BRAND / "assets/logo/wordmark-transparent.png"), width=Mm(TOKENS["minimum_logo_width_mm"]["wordmark"]))
        p.add_run("\t").add_picture(qr(), width=Mm(17))
        paragraph(cell, TOKENS["lead"], 12, bold=True, after=3)
        paragraph(cell, "Lead Developer / Product Manager", 8.5, after=5)
        p = paragraph(cell, after=3)
        link(p, TOKENS["maintainer_email"], "mailto:" + TOKENS["maintainer_email"], 8.5)
        p = paragraph(cell, after=5)
        link(p, TOKENS["domain"], "https://" + TOKENS["domain"], 8.5)
        run(p, "  /  ", 8.5, ink=MUTED)
        link(p, TOKENS["maintainer_domain"], "https://" + TOKENS["maintainer_domain"], 8.5)
        paragraph(cell, "Created & maintained by " + TOKENS["maintainer"], 8, after=0, ink=MUTED)
    return doc


def stickers():
    doc = base("FNLLA sticker sheet", top=23, side=30)
    doc.sections[0].bottom_margin = Mm(15)
    for cell in grid(doc, 70, 40, 5, 10):
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(62), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run().add_picture(str(BRAND / "assets/logo/wordmark-transparent.png"), width=Mm(TOKENS["minimum_logo_width_mm"]["wordmark"]))
        p.add_run("\t").add_picture(qr(), width=Mm(19))
        p = paragraph(cell, after=4)
        link(p, TOKENS["domain"], "https://" + TOKENS["domain"], 12)
        paragraph(cell, "Web framework. PHP foundation.", 8.5, after=3)
        paragraph(cell, "Created & maintained by " + TOKENS["maintainer"], 7.5, after=0, ink=MUTED)
    ending = paragraph(doc, after=0)
    ending.paragraph_format.line_spacing = Pt(1)
    run(ending, "", 1)
    return doc


def save(doc, name):
    """Embed full OFL fonts, so edits can introduce new characters without a subset limit."""
    buffer = BytesIO()
    doc.save(buffer)
    with ZipFile(buffer) as archive:
        parts = {item: archive.read(item) for item in archive.namelist()}
    fonts = etree.fromstring(parts["word/fontTable.xml"])
    rels = etree.Element("Relationships", nsmap={None: "http://schemas.openxmlformats.org/package/2006/relationships"})
    for family, prefix in [(FONT, "SpaceGrotesk"), (MONO, "JetBrainsMono")]:
        font = element("w:font", **{"w:name": family})
        fonts.append(font)
        for style, file in [("Regular", "Regular"), ("Bold", "SemiBold")]:
            source = BRAND / "assets/fonts" / f"{prefix}-{file}.ttf"
            key = uuid.uuid5(uuid.NAMESPACE_URL, "https://fnlla.com/fonts/" + source.name)
            data = bytearray(source.read_bytes())
            # ECMA-376 font embedding reverses the GUID bytes, then XORs its first 32 bytes.
            mask = key.bytes[::-1]
            for i in range(32):
                data[i] ^= mask[i % 16]
            part = f"fonts/{prefix}-{file}.odttf"
            parts["word/" + part] = bytes(data)
            rid = "font" + prefix + style
            etree.SubElement(rels, "Relationship", Id=rid, Type=NS["r"] + "/font", Target=part)
            font.append(element("w:embed" + style, **{"r:id": rid, "w:fontKey": "{" + str(key).upper() + "}", "w:subsetted": "false"}))
    types = etree.fromstring(parts["[Content_Types].xml"])
    etree.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Default", Extension="odttf", ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont")
    parts["[Content_Types].xml"] = etree.tostring(types)
    parts["word/fontTable.xml"] = etree.tostring(fonts)
    parts["word/_rels/fontTable.xml.rels"] = etree.tostring(rels)
    # Stable ZIP headers make rebuilds reviewable; no workstation paths enter the templates.
    OUT.mkdir(parents=True, exist_ok=True)
    with ZipFile(OUT / name, "w", ZIP_DEFLATED) as archive:
        for path, data in sorted(parts.items()):
            info = ZipInfo(path, (2026, 9, 7, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            archive.writestr(info, data)
    print(name)


if __name__ == "__main__":
    for name, doc in [("FNLLA-Cover-Letter.docx", letter(True)), ("FNLLA-Letterhead.docx", letter()),
                      ("FNLLA-Business-Cards.docx", cards()), ("FNLLA-Sticker-Sheet.docx", stickers())]:
        save(doc, name)
